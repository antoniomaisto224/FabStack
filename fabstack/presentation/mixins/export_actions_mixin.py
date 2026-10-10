"""Image, Word, and text export behavior."""

import io
import logging
import math
import os
import tempfile

import matplotlib.colors as mcolors
import matplotlib.patches as mpatches
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.offsetbox import AnnotationBbox, OffsetImage
from PyQt5.QtCore import QTimer
from PyQt5.QtGui import QPixmap
from PyQt5.QtWidgets import (
    QApplication,
    QDialog,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QLabel,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QVBoxLayout,
)
try:
    from PIL import Image
except ImportError:
    Image = None

from fabstack.config.visualization import (
    EXPORT_PREVIEW_SUBPLOTS,
    FEATURE_EMPHASIS_DEFAULT,
    FEATURE_EMPHASIS_MAX,
    FEATURE_EMPHASIS_MIN,
    PREVIEW_EXPORT_AS_MAIN_VIEW,
    VISUAL_SCALE_MAX,
    VISUAL_SCALE_MIN,
)
from fabstack.materials.chemistry import chem_subs_unicode, legend_label
from fabstack.rendering.constants import (
    DISPLAY_TOP_HEADROOM_PX,
    DYNAMIC_TARGET_STACK_PX,
    FEATURE_SCALE_COEFFICIENT,
    FEATURE_SCALE_EXPONENT,
    MAX_AUTO_CANVAS_HEIGHT_SCALE,
    MAX_CONSTRAINT_PASSES,
    MAX_THICKNESS_NM,
    MIN_LAYER_PX,
    MIN_SUBSTRATE_CORE_PX,
    MIN_THICKNESS_NM,
    MIN_VOID_PX,
    ORDER_MARGIN_PX,
    ORDER_TIE_RATIO,
    SUBSTRATE_ROWS_PER_UM,
    VISUAL_RENDER_SCALE_BASE,
    WORD_EXPORT_DPI,
    WORD_EXPORT_HORIZONTAL_PAD_IN,
    WORD_EXPORT_IMAGE_WIDTH_IN,
    WORD_EXPORT_PAGE_HEIGHT_IN,
    WORD_EXPORT_PAGE_MARGIN_IN,
    WORD_EXPORT_PAGE_WIDTH_IN,
    WORD_EXPORT_VERTICAL_PAD_PX,
)
from fabstack.rendering.scaling import (
    MAX_LAYER_PX,
    _raw_feature_scale_um,
    _visual_height_to_display_rows,
    physical_span_to_raw_visual_px,
    thickness_to_raw_visual_px,
)
from fabstack.resources import LOGO_PATH

try:
    from docx import Document
    from docx.shared import Inches
except ImportError:
    Document = None

LOGGER = logging.getLogger("fabstack.presentation.mixins.plotting_export_mixin")


class ExportActionsMixin:
    def _save_cropped_figure(self, fig, out_path):
        """
        Save `fig` as a PNG cropped to its active content (stack + legend),
        avoiding large white margins and preventing legend clipping.

        Strategy:
        - First render with bbox_inches='tight' to ensure the legend is fully
          inside the image (no clipping).
        - Then crop all near-white borders away and save the final PNG.
        """
        # Render to an in-memory buffer with tight bounding box
        buf = io.BytesIO()
        fig.savefig(buf, dpi=150, bbox_inches="tight", pad_inches=0.05)
        buf.seek(0)

        # Load image as array (H, W, 4) or (H, W, 3)
        img = plt.imread(buf)
        buf.close()

        if img.ndim == 2:
            # grayscale → fake RGB
            rgb = np.stack([img, img, img], axis=-1)
        else:
            rgb = img[:, :, :3]

        # Values are in [0,1]; treat pixels with all channels >= 0.98 as background
        h, w, _ = rgb.shape
        mask = np.any(rgb < 0.98, axis=2)

        ys = np.where(np.any(mask, axis=1))[0]
        xs = np.where(np.any(mask, axis=0))[0]

        # If nothing non-white, just save as-is
        if ys.size == 0 or xs.size == 0:
            # Re-save to the final path
            fig.savefig(out_path, dpi=150, bbox_inches="tight", pad_inches=0.05)
            return

        y0, y1 = ys[0], ys[-1]
        x0, x1 = xs[0], xs[-1]

        pad = 4  # small extra padding around content
        y0 = max(0, y0 - pad)
        y1 = min(h - 1, y1 + pad)
        x0 = max(0, x0 - pad)
        x1 = min(w - 1, x1 + pad)

        cropped = img[y0:y1 + 1, x0:x1 + 1, ...]
        plt.imsave(out_path, cropped)

    def _bond_source_export_lines(self, step: dict):
        """
        Build textual lines describing the source process of a Bond-to step.
        """
        if not isinstance(step, dict):
            return []
        src = step.get("bond_source")
        if not isinstance(src, dict):
            return []

        src_path = str(src.get("path", "")).strip()
        src_name = os.path.basename(src_path) if src_path else "another process"
        lines = [f"Bonded process: {src_name}"]

        titles = src.get("history_titles", [])
        if isinstance(titles, list) and titles:
            lines.append("Bonded process steps:")
            for i, title in enumerate(titles, start=1):
                t = str(title).strip() if title is not None else ""
                lines.append(f"  {i}. {t if t else '(untitled)'}")
        return lines

    # ---------- Export ----------
    def _show_export_status(self, message: str) -> None:
        if hasattr(self, "statusBar"):
            self.statusBar().showMessage(str(message), 4000)

    def _show_nothing_to_export(self) -> None:
        QMessageBox.information(
            self,
            "Nothing to export",
            "Add at least one process step, then try again.",
        )

    def _run_export_action(self, label: str, func):
        try:
            return func()
        except Exception:
            LOGGER.exception("%s failed", label)
            QMessageBox.critical(
                self,
                f"Couldn’t {label.lower()}",
                "Try again. If the problem continues, restart FabStack.",
            )
            return None

    def export_history_image(self):
        if not self.history:
            self._show_nothing_to_export()
            return
        fname, _ = QFileDialog.getSaveFileName(
            self,
            "Save process image",
            "",
            "PNG image (*.png)",
        )
        if not fname:
            return
        if not fname.lower().endswith(".png"):
            fname += ".png"
        n = len(self.history)
        cols = 3
        rows = (n + cols - 1) // cols
        # Grow the export canvas only above 1.0; below 1.0 the stack itself
        # is compressed inside the normal figure.
        scale = self._visual_render_scale_factor()
        canvas_scale = max(1.0, scale)
        fig, axes = plt.subplots(
            rows,
            cols,
            figsize=(6 * cols + 2, 5 * canvas_scale * rows)
        )
        axes = np.array(axes).reshape(-1)
        process_profile = self._process_display_adjustment_profile()
        for i, st in enumerate(self.history):
            self.plot_grid(
                axes[i],
                st["grid"],
                st["title"],
                main_axis=False,
                history_index=i,
                include_export_logo=True,
                row_heights_um=st.get("row_heights_um"),
                visual_geometry=st.get("visual_geometry", []),
                process_profile=process_profile,
            )
        for j in range(i + 1, len(axes)):
            axes[j].axis("off")
        fig.subplots_adjust(left=0.06, right=0.96, top=0.95, bottom=0.06, wspace=0.35, hspace=0.6)

        # Use the cropped figure saver so the canvas width dynamically grows
        # to include legends and then trims only true whitespace.
        self._save_cropped_figure(fig, fname)
        plt.close(fig)
        self._show_export_status("Process image saved.")

    def export_current_view(self):
        if self.xsec is None:
            QMessageBox.information(
                self,
                "Nothing to copy",
                "Create or open a process first.",
            )
            return

        tmp_path = os.path.join(
            tempfile.gettempdir(),
            f"fabstack_clip_{os.getpid()}.png"
        )
        logo_artist = self._add_export_logo_to_axis(self.canvas.ax)
        try:
            # Use cropped save so clipboard image has no excessive whitespace
            self._save_cropped_figure(self.canvas.fig, tmp_path)
        finally:
            if logo_artist is not None:
                try:
                    logo_artist.remove()
                    self.canvas.draw_idle()
                except Exception:
                    pass

        pixmap = QPixmap(tmp_path)
        if pixmap.isNull():
            raise RuntimeError("Could not create clipboard image from exported PNG.")

        app = QApplication.instance()
        if app is None:
            raise RuntimeError("Clipboard is unavailable because no Qt application instance exists.")
        app.clipboard().setPixmap(pixmap)

        try:
            os.remove(tmp_path)
        except Exception:
            pass

        self._show_export_status("Current view copied.")

    def export_process_word(self):
        if not self.history:
            self._show_nothing_to_export()
            return

        if Document is None:
            LOGGER.warning("Word export support is unavailable")
            QMessageBox.warning(
                self,
                "Word export unavailable",
                "Export a PNG process image instead.",
            )
            return

        fname, _ = QFileDialog.getSaveFileName(
            self,
            "Save process document",
            "",
            "Word document (*.docx)",
        )
        if not fname:
            return
        if not fname.lower().endswith(".docx"):
            fname += ".docx"

        doc = Document()
        self._configure_word_export_page(doc)
        doc.add_heading("Process steps", level=1)
        image_width = self._word_export_image_width_inches(doc)

        tmp_files = []
        rendered_steps = []
        try:
            process_profile = self._process_display_adjustment_profile()
            for i, st in enumerate(self.history):
                title = st["title"]
                grid = st["grid"]

                # Grow the export canvas only above 1.0; below 1.0 the stack
                # itself is compressed inside the normal figure.
                scale = self._visual_render_scale_factor()
                fig, ax = plt.subplots(figsize=(image_width, 4 * max(1.0, scale)))
                self.plot_grid(
                    ax,
                    grid,
                    title=None,
                    main_axis=False,
                    history_index=i,
                    include_export_logo=True,
                    row_heights_um=st.get("row_heights_um"),
                    visual_geometry=st.get("visual_geometry", []),
                    process_profile=process_profile,
                )
                fig.subplots_adjust(left=0.06, right=0.94, top=0.95, bottom=0.05)

                tmp_path = os.path.join(
                    tempfile.gettempdir(),
                    f"fabstack_step_{os.getpid()}_{i}.png"
                )

                rendered_steps.append({
                    "index": i,
                    "step": st,
                    "title": title,
                    "path": tmp_path,
                    "payload": self._render_word_step_figure_payload(
                        fig,
                        ax,
                        grid,
                        i,
                        row_heights_um=st.get("row_heights_um"),
                        process_profile=process_profile,
                    ),
                })

                plt.close(fig)
                tmp_files.append(tmp_path)

            self._save_word_step_images_with_consistent_stack_width(rendered_steps)

            for item in rendered_steps:
                st = item["step"]
                p = doc.add_paragraph()

                run = p.add_run(f"{item['index'] + 1}. {item['title']}")
                run.bold = True
                for line in self._bond_source_export_lines(st):
                    doc.add_paragraph(line)

                doc.add_picture(item["path"], width=Inches(image_width))
                doc.add_paragraph()

            doc.save(fname)
            self._show_export_status("Process document saved.")
        finally:
            for path in tmp_files:
                try:
                    if os.path.exists(path):
                        os.remove(path)
                except Exception:
                    pass

    def _word_export_image_width_inches(self, doc):
        try:
            section = doc.sections[0]
            usable = (
                float(section.page_width)
                - float(section.left_margin)
                - float(section.right_margin)
            ) / 914400.0
            return min(WORD_EXPORT_IMAGE_WIDTH_IN, max(1.0, usable))
        except Exception:
            return WORD_EXPORT_IMAGE_WIDTH_IN

    def _configure_word_export_page(self, doc):
        try:
            section = doc.sections[0]
            section.page_width = Inches(WORD_EXPORT_PAGE_WIDTH_IN)
            section.page_height = Inches(WORD_EXPORT_PAGE_HEIGHT_IN)
            margin = Inches(WORD_EXPORT_PAGE_MARGIN_IN)
            section.left_margin = margin
            section.right_margin = margin
            section.top_margin = margin
            section.bottom_margin = margin
        except Exception:
            pass

    def _image_array_to_uint8_rgb(self, img):
        arr = np.asarray(img)
        if arr.ndim == 2:
            arr = np.stack([arr, arr, arr], axis=-1)

        if arr.dtype.kind == "f":
            arr = np.clip(arr, 0.0, 1.0)
            arr = np.rint(arr * 255.0).astype(np.uint8)
        elif arr.dtype != np.uint8:
            arr = np.clip(arr, 0, 255).astype(np.uint8)

        if arr.ndim != 3:
            return np.full((1, 1, 3), 255, dtype=np.uint8)

        if arr.shape[2] >= 4:
            rgb = arr[:, :, :3].astype(float)
            alpha = arr[:, :, 3:4].astype(float) / 255.0
            arr = np.rint(rgb * alpha + 255.0 * (1.0 - alpha)).astype(np.uint8)
        elif arr.shape[2] == 1:
            arr = np.repeat(arr, 3, axis=2)
        else:
            arr = arr[:, :, :3]

        return arr

    def _save_uint8_rgb_image(self, out_path, image):
        image = np.asarray(image, dtype=np.uint8)
        if Image is not None:
            Image.fromarray(image).save(out_path)
        else:
            plt.imsave(out_path, image)

    def _resize_uint8_rgb_width(self, image, target_width):
        image = np.asarray(image, dtype=np.uint8)
        target_width = max(1, int(round(target_width)))
        if image.ndim != 3 or image.shape[1] == target_width:
            return image.copy()

        if Image is not None:
            resample = getattr(getattr(Image, "Resampling", Image), "LANCZOS", Image.BICUBIC)
            resized = Image.fromarray(image).resize(
                (target_width, int(image.shape[0])),
                resample=resample,
            )
            return np.asarray(resized, dtype=np.uint8)

        indices = np.linspace(0, image.shape[1] - 1, target_width)
        indices = np.clip(np.rint(indices).astype(int), 0, image.shape[1] - 1)
        return image[:, indices, :].copy()

    def _render_word_step_figure_payload(
        self,
        fig,
        ax,
        grid,
        history_index,
        row_heights_um=None,
        process_profile=None,
    ):
        buf = io.BytesIO()
        fig.set_dpi(WORD_EXPORT_DPI)
        fig.canvas.draw()
        try:
            renderer = fig.canvas.get_renderer()
            tight_bbox = fig.get_tightbbox(renderer).padded(WORD_EXPORT_HORIZONTAL_PAD_IN)
        except Exception:
            tight_bbox = None

        save_kwargs = {
            "dpi": WORD_EXPORT_DPI,
            "facecolor": "white",
            "edgecolor": "white",
        }
        if tight_bbox is not None:
            save_kwargs["bbox_inches"] = tight_bbox
        fig.savefig(buf, **save_kwargs)
        buf.seek(0)
        img = self._image_array_to_uint8_rgb(plt.imread(buf))
        buf.close()

        display_grid = self._display_grid_auto(
            grid,
            row_heights_um=(
                row_heights_um
                if row_heights_um is not None
                else self._row_heights_for_plot_grid(grid, history_index)
            ),
            history_index=history_index,
            process_profile=process_profile,
        )
        try:
            stack_w_data = int(np.asarray(display_grid).shape[1])
        except Exception:
            stack_w_data = 0

        stack_x0 = None
        stack_x1 = None
        if ax is not None and stack_w_data > 0:
            try:
                p0 = ax.transData.transform((-0.5, 0.0))
                p1 = ax.transData.transform((float(stack_w_data) - 0.5, 0.0))
                x0, x1 = sorted((float(p0[0]), float(p1[0])))
                if tight_bbox is not None:
                    bbox_x0_px = float(tight_bbox.x0) * float(WORD_EXPORT_DPI)
                    bbox_w_px = max(1.0, float(tight_bbox.width) * float(WORD_EXPORT_DPI))
                    scale_x = float(img.shape[1]) / bbox_w_px
                    stack_x0 = (x0 - bbox_x0_px) * scale_x
                    stack_x1 = (x1 - bbox_x0_px) * scale_x
                else:
                    bbox_w = float(getattr(fig.bbox, "width", img.shape[1])) or float(img.shape[1])
                    scale_x = float(img.shape[1]) / bbox_w
                    stack_x0 = x0 * scale_x
                    stack_x1 = x1 * scale_x
                stack_x0 = max(0.0, min(float(img.shape[1]), stack_x0))
                stack_x1 = max(0.0, min(float(img.shape[1]), stack_x1))
            except Exception:
                stack_x0 = None
                stack_x1 = None

        if stack_x0 is None or stack_x1 is None or stack_x1 <= stack_x0:
            stack_x0 = 0.0
            stack_x1 = float(img.shape[1])

        mask = np.any(img[:, :, :3] < 250, axis=2)
        ys = np.where(np.any(mask, axis=1))[0]
        if ys.size == 0:
            return {
                "image": img,
                "stack_x0": stack_x0,
                "stack_x1": stack_x1,
            }

        pad = WORD_EXPORT_VERTICAL_PAD_PX
        y0 = max(0, int(ys[0]) - pad)
        y1 = min(img.shape[0] - 1, int(ys[-1]) + pad)
        cropped = img[y0:y1 + 1, :, ...]
        return {
            "image": cropped,
            "stack_x0": stack_x0,
            "stack_x1": stack_x1,
        }

    def _save_word_step_images_with_consistent_stack_width(self, rendered_steps):
        payloads = [
            item.get("payload")
            for item in rendered_steps
            if isinstance(item.get("payload"), dict)
        ]
        if not payloads:
            return

        stack_widths = [
            max(1.0, float(payload.get("stack_x1", 0.0)) - float(payload.get("stack_x0", 0.0)))
            for payload in payloads
        ]
        target_stack_width = max(1, int(round(max(stack_widths))))

        prepared = []
        for item in rendered_steps:
            payload = item.get("payload") if isinstance(item.get("payload"), dict) else {}
            image = self._image_array_to_uint8_rgb(payload.get("image", np.full((1, 1, 3), 255, dtype=np.uint8)))
            stack_x0 = float(payload.get("stack_x0", 0.0))
            stack_x1 = float(payload.get("stack_x1", image.shape[1]))
            current_stack_width = max(1.0, stack_x1 - stack_x0)
            scale_x = float(target_stack_width) / current_stack_width

            if abs(scale_x - 1.0) > 0.002:
                new_width = max(1, int(round(float(image.shape[1]) * scale_x)))
                image = self._resize_uint8_rgb_width(image, new_width)
                stack_x0 *= scale_x
                stack_x1 *= scale_x

            prepared.append({
                "item": item,
                "image": image,
                "stack_x0": stack_x0,
                "stack_x1": stack_x1,
            })

        target_stack_left = int(round(max(float(entry["stack_x0"]) for entry in prepared)))
        final_width = int(round(max(
            target_stack_left + (float(entry["image"].shape[1]) - float(entry["stack_x0"]))
            for entry in prepared
        )))
        final_width = max(final_width, target_stack_left + target_stack_width)

        for entry in prepared:
            image = entry["image"]
            offset_x = int(round(float(target_stack_left) - float(entry["stack_x0"])))
            height = int(image.shape[0])
            canvas = np.full((height, final_width, 3), 255, dtype=np.uint8)

            dst_x0 = max(0, offset_x)
            src_x0 = max(0, -offset_x)
            copy_w = min(int(image.shape[1]) - src_x0, final_width - dst_x0)
            if copy_w > 0:
                canvas[:, dst_x0:dst_x0 + copy_w, :] = image[:, src_x0:src_x0 + copy_w, :]

            self._save_uint8_rgb_image(entry["item"]["path"], canvas)

    def _save_word_step_figure(self, fig, out_path):
        payload = self._render_word_step_figure_payload(fig, None, np.empty((1, 1)), -1)
        self._save_uint8_rgb_image(out_path, payload["image"])

    def export_step_list_txt(self):
        """
        Export only the list of step titles (no images, no logo)
        as a plain text file.
        """
        if not self.history:
            self._show_nothing_to_export()
            return

        fname, _ = QFileDialog.getSaveFileName(
            self,
            "Save step list",
            "",
            "Text file (*.txt);;All files (*)",
        )
        if not fname:
            return
        root, ext = os.path.splitext(fname)
        if not ext:
            fname = root + ".txt"

        try:
            with open(fname, "w", encoding="utf-8") as f:
                for i, st in enumerate(self.history, start=1):
                    title = st.get("title", "")
                    f.write(f"{i}. {title}\n")
                    for line in self._bond_source_export_lines(st):
                        f.write(f"   {line}\n")
            self._show_export_status("Step list saved.")
        except Exception:
            LOGGER.exception("Could not save text step list to %s", fname)
            QMessageBox.warning(
                self,
                "Export failed",
                "The step list could not be saved. Choose a different location and try again.",
            )

    def export_step_list_word(self):
        """
        Export only the list of step titles as a Word document
        (no images, no logo), one line per step.
        """
        if not self.history:
            self._show_nothing_to_export()
            return

        if Document is None:
            LOGGER.warning("Word export support is unavailable")
            QMessageBox.warning(
                self,
                "Word export unavailable",
                "Export the step list as TXT instead.",
            )
            return

        fname, _ = QFileDialog.getSaveFileName(
            self,
            "Save step list",
            "",
            "Word document (*.docx)",
        )
        if not fname:
            return
        if not fname.lower().endswith(".docx"):
            fname += ".docx"

        try:
            doc = Document()
            doc.add_heading("Process step list", level=1)

            for i, st in enumerate(self.history, start=1):
                title = st.get("title", "")
                p = doc.add_paragraph()
                p.add_run(f"{i}. {title}")
                for line in self._bond_source_export_lines(st):
                    doc.add_paragraph(line)

            doc.save(fname)
            self._show_export_status("Step list saved.")
        except Exception:
            LOGGER.exception("Could not save Word step list to %s", fname)
            QMessageBox.warning(
                self,
                "Export failed",
                "The step list could not be saved. Choose a different location and try again.",
            )

    # -- extracted methods --
